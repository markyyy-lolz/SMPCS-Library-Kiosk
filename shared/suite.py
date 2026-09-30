"""Library suite services and validated import/export helpers."""
import base64
import csv
import io
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from shared.api import ApiError

MIGRATION_HELP='Apply migrations/003_library_suite.sql in Supabase SQL Editor, then sign in again.'


def suite(api, user, action, data=None):
    try:
        return api.rpc('library_suite', {'p_token':user.get('token',''), 'p_action':action, 'p_data':data or {}})
    except ApiError as exc:
        if 'PGRST202' in str(exc) or 'Could not find the function' in str(exc):
            raise ApiError(MIGRATION_HELP) from exc
        raise


def station_suite(host, action, data=None):
    return host.api.rpc('library_suite_station', {'p_station':host.cfg['STATION_CODE'], 'p_token':host.cfg['STATION_TOKEN'], 'p_action':action, 'p_data':data or {}})


class StaffAPI:
    """Keep credentials private and route all staff writes through live role checks."""
    def __init__(self, api, user): self.api, self.user = api, user
    def __getattr__(self, name): return getattr(self.api, name)
    def rpc(self, name, params=None):
        if name in ('library_verify_loan','library_reject_loan'):
            from shared.operations import operations
            return operations(self.api,self.user,'verify_loan' if name=='library_verify_loan' else 'reject_loan',{'id':(params or {}).get('p_loan_id')})
        if name in ('library_save_book','library_register_station'):
            return self.api.rpc('library_staff_rpc', {'p_token':self.user.get('token',''), 'p_name':name, 'p_args':params or {}})
        return self.api.rpc(name, params)


def read_import(path, kind):
    fields={'members':['member_no','rfid_uid','full_name','member_type','grade_level','section'],
            'books':['accession_no','book_rfid','title','author','category','shelf','total_copies']}[kind]
    if Path(path).suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            iterator=book.active.iter_rows(values_only=True);headers=[str(x or '').strip() for x in next(iterator)]
            rows=[dict(zip(headers,r)) for r in iterator if any(v is not None for v in r)]
        finally:book.close()
    else:
        with open(path,encoding='utf-8-sig',newline='') as f:
            reader=csv.DictReader(f);headers=reader.fieldnames or [];rows=list(reader)
    required=fields[:3] if kind=='members' else [fields[0],fields[2]]
    required=[f for f in required if f!='rfid_uid']
    if not set(required)<=set(headers):raise ValueError('Required columns: '+', '.join(required))
    if not rows or len(rows)>1000:raise ValueError('Import 1–1000 rows per file.')
    result=[];seen={fields[0]:set(),fields[1]:set()}
    for n,row in enumerate(rows,2):
        item={k:str(row.get(k) if row.get(k) is not None else '').strip() for k in fields}
        if any(not item[k] for k in required):raise ValueError(f'Row {n}: required value missing.')
        for k,values in seen.items():
            if item[k] and item[k] in values:raise ValueError(f'Row {n}: duplicate {k}: {item[k]}')
            if item[k]:values.add(item[k])
        if kind=='books':
            try:item['total_copies']=int(item['total_copies'] or 1)
            except ValueError:raise ValueError(f'Row {n}: copies must be a whole number.')
            if item['total_copies']<1:raise ValueError(f'Row {n}: copies must be positive.')
        result.append(item)
    return fields,result


def due_hint(value, status, today=None):
    if status!='borrowed':return '—'
    try:date=datetime.fromisoformat(value.replace('Z','+00:00')).astimezone(timezone(timedelta(hours=8))).date()
    except (ValueError,AttributeError):return '—'
    delta=(date-(today or datetime.now(timezone(timedelta(hours=8))).date())).days
    return f'{-delta} days overdue' if delta<0 else 'Due today' if delta==0 else 'Due tomorrow' if delta==1 else f'Due in {delta} days'


def export_rows(path, headers, rows, title='Library report'):
    from shared.admin_features import csv_cell
    suffix=Path(path).suffix.lower()
    if suffix=='.xlsx':
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.title='Report';sheet.append(headers)
        for row in rows:sheet.append([csv_cell(v) for v in row])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(45,max(16,max(len(str(c.value or '')) for c in col)+2))
        book.save(path)
    elif suffix=='.pdf':
        from xml.sax.saxutils import escape
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import A4, landscape
        style=getSampleStyleSheet();page=landscape(A4);doc=SimpleDocTemplate(str(path),pagesize=page,rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
        body=style['BodyText'];body.fontSize=8;body.leading=11
        values=[[Paragraph(escape(str(v or '')),body) for v in row] for row in [headers]+rows]
        table=Table(values,colWidths=[(page[0]-48)/len(headers)]*len(headers),repeatRows=1)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DBEAFE')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),8),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#DCE5EF'))]))
        doc.build([Paragraph(escape(title),style['Title']),Spacer(1,12),table])
    else:
        with open(path,'w',encoding='utf-8-sig',newline='') as f:
            writer=csv.writer(f);writer.writerow(headers);writer.writerows([[csv_cell(v) for v in row] for row in rows])


def diagnostics():
    # Construct an allowlist report: never include configs, credentials, or user records.
    import platform
    from shared.updates import VERSION
    from shared.config import APP_DIR
    logs=APP_DIR/'logs'
    return json.dumps({'app':'SMPCS Library','version':VERSION,'os':platform.platform(),
        'python':platform.python_version(),'generated_at':datetime.now(timezone.utc).isoformat(),
        'log_files':[{'name':p.name,'bytes':p.stat().st_size} for p in logs.glob('*.log')]},indent=2)


def cover_image(path):
    from PIL import Image
    with Image.open(path) as image:
        image.thumbnail((360,480)); image=image.convert('RGB');out=io.BytesIO();image.save(out,format='JPEG',quality=80)
    return base64.b64encode(out.getvalue()).decode()


def error_report(message):
    report=json.loads(diagnostics())
    codes=re.findall(r'\b(?:PGRST\d{3}|HTTP [45]\d\d|[45]\d{4})\b',str(message))
    report['error_codes']=codes
    report['category']='network' if any(x in str(message).lower() for x in ('network','timeout','connect')) else 'application/database'
    return json.dumps(report,indent=2)


def add_copy_error(box,message):
    from PyQt6.QtWidgets import QMessageBox,QApplication
    b=box.addButton('Copy error details',QMessageBox.ButtonRole.ActionRole)
    b.clicked.connect(lambda:QApplication.clipboard().setText(error_report(message)))
